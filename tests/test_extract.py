import io

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter

from app.extract import MAX_PHOTOS, extract_text, prepare_upload

TEXT = "The sun is a star at the center of our solar system. " * 2


def make_docx(text):
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def make_blank_pdf():  # a PDF with a page but no text inside, like a scan
    buf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buf)
    return buf.getvalue()


def make_pptx(text):
    buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = text
    prs.save(buf)
    return buf.getvalue()


def test_extract_txt_and_md():
    assert extract_text("a.txt", b"hello") == "hello"
    assert extract_text("a.MD", b"# hi") == "# hi"


def test_extract_docx():
    assert "sun" in extract_text("a.docx", make_docx(TEXT))


def test_extract_pptx():
    assert "sun" in extract_text("a.pptx", make_pptx(TEXT))


def test_extract_docx_includes_tables():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Q1 what is x"
    doc.save(buf)
    assert "Q1 what is x" in extract_text("a.docx", buf.getvalue())


def test_extract_unsupported_type():
    with pytest.raises(ValueError):
        extract_text("a.exe", b"x")


# --- prepare_upload: text or the files themselves ---

def test_prepare_document_gives_text():
    assert "sun" in prepare_upload([("notes.docx", make_docx(TEXT))])


def test_prepare_photos_gives_files():
    files = [("page1.JPG", b"img1"), ("page2.heic", b"img2")]
    assert prepare_upload(files) == [(b"img1", "image/jpeg"), (b"img2", "image/heic")]


def test_prepare_scanned_pdf_gives_the_pdf():
    pdf = make_blank_pdf()
    assert prepare_upload([("scan.pdf", pdf)]) == [(pdf, "application/pdf")]


def test_prepare_too_many_photos():
    with pytest.raises(ValueError, match="at most"):
        prepare_upload([(f"p{i}.png", b"x") for i in range(MAX_PHOTOS + 1)])


@pytest.mark.parametrize("files", [
    [("page.jpg", b"x"), ("notes.docx", b"x")],  # photo mixed with a document
    [("a.txt", b"x"), ("b.txt", b"x")],  # two documents
])
def test_prepare_refuses_mixed_or_several_documents(files):
    with pytest.raises(ValueError, match="one document"):
        prepare_upload(files)


def test_prepare_no_files():
    with pytest.raises(ValueError):
        prepare_upload([])