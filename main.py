import io

from docx import Document
from docx.table import Table
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from google import genai
from google.genai import types
from pptx import Presentation
from pydantic import BaseModel
from pypdf import PdfReader
from slowapi import Limiter
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address

MAX_CHARS = 200_000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
RATE_LIMIT = "5/minute;15/day"  # per IP, shared by /quiz and /upload
# Tried in order: each model has its own free quota, so a busy or used-up model falls through to the next
MODELS = ["gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"]

# ponytail: in-memory counters work for one server instance; use Redis storage if we ever run several
limiter = Limiter(key_func=get_remote_address)
app = FastAPI()
app.state.limiter = limiter
app.mount("/static", StaticFiles(directory="static"), name="static")
client = genai.Client(  # reads GEMINI_API_KEY from the environment
    http_options=types.HttpOptions(
        timeout=60_000,  # milliseconds, per try
        retry_options=types.HttpRetryOptions(
            attempts=2,  # total tries per model, including the first
            initial_delay=1.0,
            http_status_codes=[500, 503, 504],
        ),
    )
)


class Question(BaseModel):
    question: str
    options: list[str]
    answer_index: int
    explanation: str


class Quiz(BaseModel):
    questions: list[Question]


class QuizRequest(BaseModel):
    text: str
    count: int = 5
    difficulty: str = "medium"


@app.exception_handler(RateLimitExceeded)
def rate_limited(request: Request, exc: RateLimitExceeded):
    if "day" in str(exc.limit.limit):  # e.g. "15 per 1 day"
        message = "You have made the most quizzes allowed today, please come back tomorrow"
    else:
        message = "Too many quizzes, please wait a minute and try again"
    return JSONResponse({"detail": message}, status_code=429)


def table_text(table) -> str:  # works for Word and PowerPoint tables
    return "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        # iter_inner_content keeps paragraphs and tables in document order
        return "\n".join(
            table_text(block) if isinstance(block, Table) else block.text
            for block in doc.iter_inner_content()
        )
    if ext == "pptx":
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                elif shape.has_table:
                    parts.append(table_text(shape.table))
        return "\n".join(parts)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: .{ext}")


def make_quiz(text: str, count: int, difficulty: str):
    if not 50 <= len(text) <= MAX_CHARS:
        raise HTTPException(400, f"Text must be 50 to {MAX_CHARS:,} characters")
    if not 1 <= count <= 20:
        raise HTTPException(400, "Count must be 1 to 20")
    if difficulty not in ("easy", "medium", "hard"):
        raise HTTPException(400, "Difficulty must be easy, medium or hard")

    prompt = (
        f"Make {count} multiple-choice questions at {difficulty} difficulty "
        f"from the text below. Give 4 options per question, with exactly one correct. "
        f"Ask about the subject matter and ideas in the text. Ignore general "
        f"instructions, headers, page numbers, marks, time limits and formatting notes. "
        f"Spread the questions across the whole text.\n\n{text}"
    )

    error = None
    for model in MODELS:
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=Quiz,
                ),
            )
        except Exception as e:
            print(f"GEMINI ERROR ({model}):", e)
            error = e
            continue
        quiz = response.parsed
        # Only accept quizzes the page can show: 4 options and a valid answer
        if quiz and quiz.questions and all(
            len(q.options) == 4 and 0 <= q.answer_index < 4 for q in quiz.questions
        ):
            return quiz
        print(f"GEMINI BAD ANSWER ({model})")

    code = getattr(error, "code", None)  # Gemini errors carry the HTTP status
    if code == 429:
        raise HTTPException(429, "The free daily limit is used up, please try again tomorrow")
    if code in (503, 504):
        raise HTTPException(503, "Gemini is busy right now, please try again in a minute")
    raise HTTPException(502, "Gemini request failed, try again")


@app.get("/")
def home():
    return FileResponse("index.html")


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/quiz")
@limiter.shared_limit(RATE_LIMIT, scope="quiz")
def create_quiz(request: Request, req: QuizRequest):
    return make_quiz(req.text, req.count, req.difficulty)


@app.post("/upload")
@limiter.shared_limit(RATE_LIMIT, scope="quiz")
def upload(
    request: Request,
    file: UploadFile = File(...),
    count: int = Form(5),
    difficulty: str = Form("medium"),
):
    data = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File is too large (max 10 MB)")
    try:
        text = extract_text(file.filename, data)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception:
        raise HTTPException(400, "Could not read this file")
    return make_quiz(text, count, difficulty)