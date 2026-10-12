import io
import os
from types import SimpleNamespace

os.environ.setdefault("GEMINI_API_KEY", "test-key")  # main.py creates the Gemini client on import

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

import main

api = TestClient(main.app)
TEXT = "The sun is a star at the center of our solar system. " * 2
FAKE_QUIZ = main.Quiz(
    questions=[
        main.Question(
            question="What is the sun?",
            options=["A star", "A planet", "A moon", "A comet"],
            answer_index=0,
            explanation="The sun is a star.",
        )
    ]
)


def make_docx(text):
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def make_pptx(text):
    buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = text
    prs.save(buf)
    return buf.getvalue()


def upload(name, data, **form):
    fields = {"count": "3", "difficulty": "easy", **form}
    return api.post("/upload", files={"file": (name, data)}, data=fields)


@pytest.fixture
def fake_gemini(monkeypatch):
    fake = SimpleNamespace(
        models=SimpleNamespace(generate_content=lambda **kw: SimpleNamespace(parsed=FAKE_QUIZ))
    )
    monkeypatch.setattr(main, "client", fake)


@pytest.fixture
def broken_gemini(monkeypatch):
    def boom(**kw):
        raise RuntimeError("Gemini is down")

    monkeypatch.setattr(main, "client", SimpleNamespace(models=SimpleNamespace(generate_content=boom)))

@pytest.fixture(autouse=True)
def reset_limiter():
    main.limiter.reset()

# --- extract_text ---

def test_extract_txt_and_md():
    assert main.extract_text("a.txt", b"hello") == "hello"
    assert main.extract_text("a.MD", b"# hi") == "# hi"


def test_extract_docx():
    assert "sun" in main.extract_text("a.docx", make_docx(TEXT))


def test_extract_pptx():
    assert "sun" in main.extract_text("a.pptx", make_pptx(TEXT))


def test_extract_unsupported_type():
    with pytest.raises(ValueError):
        main.extract_text("a.exe", b"x")


# --- /quiz ---

def test_home_page():
    assert api.get("/").status_code == 200


def test_quiz_ok(fake_gemini):
    r = api.post("/quiz", json={"text": TEXT, "count": 1, "difficulty": "easy"})
    assert r.status_code == 200
    assert r.json()["questions"][0]["answer_index"] == 0


def test_quiz_text_too_short():
    assert api.post("/quiz", json={"text": "short"}).status_code == 400


def test_quiz_text_too_long():
    r = api.post("/quiz", json={"text": "a" * (main.MAX_CHARS + 1)})
    assert r.status_code == 400


@pytest.mark.parametrize("count", [0, 21])
def test_quiz_bad_count(count):
    assert api.post("/quiz", json={"text": TEXT, "count": count}).status_code == 400


def test_quiz_bad_difficulty():
    assert api.post("/quiz", json={"text": TEXT, "difficulty": "insane"}).status_code == 400


def test_quiz_gemini_failure(broken_gemini):
    assert api.post("/quiz", json={"text": TEXT}).status_code == 502


# --- /upload ---

def test_upload_docx_ok(fake_gemini):
    r = upload("a.docx", make_docx(TEXT))
    assert r.status_code == 200
    assert r.json()["questions"][0]["question"] == "What is the sun?"


def test_upload_unsupported_type():
    assert upload("a.exe", b"x").status_code == 400


def test_upload_unreadable_file():
    assert upload("a.pdf", b"not really a pdf").status_code == 400


def test_upload_too_large():
    assert upload("big.txt", b"a" * (main.MAX_UPLOAD_BYTES + 1)).status_code == 413

# --- health and rate limit ---

def test_health():
    assert api.get("/health").json() == {"status": "ok"}


def test_rate_limit(fake_gemini):
    n = int(main.RATE_LIMIT.split("/")[0])
    responses = [api.post("/quiz", json={"text": TEXT}) for _ in range(n + 1)]
    assert [r.status_code for r in responses] == [200] * n + [429]
    assert "wait a minute" in responses[-1].json()["detail"]