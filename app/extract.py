"""Turns uploaded files into what we send to Gemini: plain text, or the files themselves for photos and scans."""
import io

from docx import Document
from docx.table import Table
from pptx import Presentation
from pypdf import PdfReader

# Photo types Gemini can read, with the MIME type it needs to be told
IMAGE_TYPES = {
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "webp": "image/webp",
    "heic": "image/heic",
}
MAX_PHOTOS = 10
MIN_PDF_TEXT = 50  # a PDF with fewer characters of text than this is treated as a scan


def file_ext(filename: str) -> str:
    return filename.lower().rsplit(".", 1)[-1]


def table_text(table) -> str:  # works for Word and PowerPoint tables
    return "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)


def extract_text(filename: str, data: bytes) -> str:
    ext = file_ext(filename)
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        # iter_inner_content keeps paragraphs and tables in document order
        return "\n".join(
            table_text(block) if isinstance(block, Table) else block.text
            for block in doc.iter_inner_content()
        )
    if ext == "pptx":
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                elif shape.has_table:
                    parts.append(table_text(shape.table))
        return "\n".join(parts)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: .{ext}")


def prepare_upload(files: list[tuple[str, bytes]]) -> str | list[tuple[bytes, str]]:
    """Returns the text to send to Gemini, or a list of (file bytes, MIME type) for photos and scanned PDFs."""
    if not files:
        raise ValueError("Please choose a file")
    exts = [file_ext(name) for name, _ in files]

    if all(ext in IMAGE_TYPES for ext in exts):
        if len(files) > MAX_PHOTOS:
            raise ValueError(f"Upload at most {MAX_PHOTOS} photos at once")
        return [(data, IMAGE_TYPES[ext]) for (_, data), ext in zip(files, exts)]

    if len(files) > 1:
        raise ValueError(f"Upload one document, or up to {MAX_PHOTOS} photos")

    name, data = files[0]
    text = extract_text(name, data)
    if exts[0] == "pdf" and len(text.strip()) < MIN_PDF_TEXT:
        return [(data, "application/pdf")]  # a scan: let Gemini read the pages itself
    return text